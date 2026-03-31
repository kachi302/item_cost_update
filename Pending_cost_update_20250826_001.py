import sys
import os
import time
import datetime as dt
import pyautogui
from openpyxl import load_workbook

# ===== 기본 설정 =====
pyautogui.FAILSAFE = False  # 필요시 True 권장
print('Item Cost Update Process, 5초 후 시작됩니다. 해당 화면을 다른 모니터로 이동하세요.')
time.sleep(5)

BASE = os.path.dirname(os.path.abspath(__file__))
IMG = lambda name: os.path.join(BASE, 'images', name)
EXCEL_IN = os.path.join(BASE, 'data', 'pending.xlsx')
EXCEL_OUT = os.path.join(BASE, 'data', 'Pending_JOB.xlsx')

def fmt(v):
    """UI 입력용 문자열 변환."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()

def locate(image_path, confidence=0.9, timeout=12, grayscale=False):
    """이미지를 timeout까지 기다렸다가 중앙 좌표 반환. 없으면 None."""
    end = time.time() + timeout
    while time.time() < end:
        try:
            box = pyautogui.locateOnScreen(image_path, confidence=confidence, grayscale=grayscale)
        except TypeError:
            # opencv 미설치 시 confidence 옵션이 없으므로 재시도
            box = pyautogui.locateOnScreen(image_path, grayscale=grayscale)
        if box:
            return pyautogui.center(box)
        time.sleep(0.3)
    return None

try:
    # 포커스 + 창 최대화(이미지 기준)
    print('Start - 화면 크기 max 로')
    pyautogui.click(1000, 800)
    screen_max_ce = locate(IMG('screen_max.png'), confidence=0.9, timeout=5)
    if screen_max_ce:
        pyautogui.click(screen_max_ce.x, screen_max_ce.y - 20, interval=1)
        print('Max Screen')
    else:
        print('Screen Max not found')

    # 엑셀 로드
    print('Excel File 읽기')
    wb = load_workbook(EXCEL_IN, data_only=True)
    if 'Sheet1' not in wb.sheetnames:
        raise RuntimeError("엑셀에 'Sheet1' 시트가 없습니다.")
    ws = wb['Sheet1']

    start_all = time.time()

    # 대기시간 설정
    wait_raw = pyautogui.prompt('대기 초는 얼마나 할까요? 기본은 0.5 이상을 입력하세요.')  # type: ignore
    try:
        wait_time = float(wait_raw) if wait_raw is not None else 0.5
    except ValueError:
        wait_time = 0.5
    if wait_time <= 0.4:
        pyautogui.alert('0.4 이하를 입력하였습니다. 프로그램을 다시 실행해주세요.')  # type: ignore
        raise SystemExit
    pyautogui.PAUSE = wait_time
    print('대기 시간:', wait_time)

    job = pyautogui.confirm(text='Pending Cost Update 작업을 진행하시겠습니까?', buttons=['OK', 'Cancel'])  # type: ignore
    if job != 'OK':
        print('사용자 취소')
        raise SystemExit

    # ===== 공통 상수 =====
    cost_element1 = 'Material'
    cost_element2 = 'Material Overhead'
    sub_element2 = 'Freight'

    # ===== 1) Pending Cost 입력/수정 루프 =====
    for row_ix, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        item, cost_type, cost, overhead = (fmt(row[0]), fmt(row[1]), fmt(row[2]), fmt(row[3]))
        t0 = time.time()

        ic1 = locate(IMG('item_cost1.png'), timeout=10)
        if not ic1:
            print(item, "item_cost1.png Not Found!")
            continue

        pyautogui.doubleClick(ic1)
        pyautogui.write(item, interval=0.1)
        pyautogui.press('tab')
        pyautogui.write(cost_type, interval=0.1)
        pyautogui.press('tab')

        fi = locate(IMG('find.png'), timeout=6)
        if not fi:
            print(item, 'Find 버튼 이미지 없음')
            continue
        pyautogui.click(fi)
        time.sleep(1)

        # Pending cost type 없음?
        if locate(IMG('cost_type_not_found.png'), timeout=2):
            new_btn = locate(IMG('new.png'), timeout=6)
            if not new_btn:
                print(item, 'New 버튼 이미지 없음')
                continue
            pyautogui.click(new_btn)
            pyautogui.write(item, interval=0.1)
            pyautogui.press('tab')
            pyautogui.write(cost_type, interval=0.1)

            inv_asset = locate(IMG('inv_asset.png'), timeout=3)
            if inv_asset:  # Inventory Asset
                cost_btn = locate(IMG('costs_btn.png'), timeout=6)
                if not cost_btn:
                    print(item, 'Costs 버튼 이미지 없음')
                    continue
                pyautogui.click(cost_btn)
                pyautogui.write(cost_element1, interval=0.1); pyautogui.press('tab')
                pyautogui.write(cost_element1, interval=0.1)
                pyautogui.press('tab', presses=3, interval=0.3)
                pyautogui.write(cost, interval=0.1)
                pyautogui.press('tab'); pyautogui.press('tab')
                pyautogui.write(cost_element2, interval=0.1); pyautogui.press('tab')
                pyautogui.write(sub_element2, interval=0.1); pyautogui.press('tab')
                pyautogui.press('tab', presses=2, interval=0.25)
                pyautogui.write(overhead, interval=0.1); pyautogui.press('tab')
                pyautogui.hotkey('ctrl', 's'); pyautogui.hotkey('ctrl', 'f4'); pyautogui.hotkey('ctrl', 'f4')
                ws[f'E{row_ix}'] = 'Inventory Asset Completed'
                print(item, 'Inventory Asset | 신규 생성', 'sec:', round(time.time() - t0, 2))
            else:
                print(item, 'Inventory Asset 아님')
                ws[f'E{row_ix}'] = 'Inventory Non Asset item'
        else:
            # Pending cost type 있음
            inv_asset = locate(IMG('inv_asset.png'), timeout=3)
            if inv_asset:
                cost_btn = locate(IMG('costs_btn.png'), timeout=6)
                if not cost_btn:
                    print(item, 'Costs 버튼 이미지 없음')
                    continue
                pyautogui.click(cost_btn)
                pyautogui.write(cost_element1, interval=0.1); pyautogui.press('tab')
                pyautogui.write(cost_element1, interval=0.1)
                pyautogui.press('tab'); pyautogui.press('tab')
                pyautogui.write('Item', interval=0.1); pyautogui.press('tab')
                pyautogui.write(cost, interval=0.1)
                pyautogui.press('tab'); pyautogui.press('tab')
                pyautogui.write(cost_element2, interval=0.1); pyautogui.press('tab')
                pyautogui.write(sub_element2, interval=0.1); pyautogui.press('tab')
                pyautogui.press('tab'); pyautogui.write('Total Value', interval=0.1)
                pyautogui.press('tab'); pyautogui.write(overhead, interval=0.1); pyautogui.press('tab')
                pyautogui.hotkey('ctrl', 's'); pyautogui.hotkey('ctrl', 'f4'); pyautogui.hotkey('ctrl', 'f4')
                ws[f'E{row_ix}'] = 'Inventory Asset Completed'
                print(item, 'Pending 존재 | Inventory Asset', 'sec:', round(time.time() - t0, 2))
            else:
                # 기본값/롤업 체크 해제 처리
                ri = locate(IMG('basedroll1.png'), timeout=3)
                if ri:
                    # default / based on rollup 체크 해제
                    pyautogui.click(ri.x - 130, ri.y)
                    pyautogui.click(ri.x + 130, ri.y)
                    pyautogui.press('enter'); pyautogui.hotkey('ctrl', 's'); pyautogui.press('tab')
                else:
                    ri_undefault = locate(IMG('item_cost_not_default.png'), timeout=3)
                    if ri_undefault:
                        pyautogui.click(ri_undefault.x + 115, ri_undefault.y + 10)
                        pyautogui.press('enter'); pyautogui.hotkey('ctrl', 's'); pyautogui.press('tab')
                        ws[f'E{row_ix}'] = 'not default, inventory asset , based on rollup'

                # 이후 동일하게 비용 입력
                cost_btn = locate(IMG('costs_btn.png'), timeout=6)
                if not cost_btn:
                    print(item, 'Costs 버튼 이미지 없음')
                    continue
                pyautogui.click(cost_btn)
                pyautogui.write(cost_element1, interval=0.1); pyautogui.press('tab')
                pyautogui.write(cost_element1, interval=0.1)
                pyautogui.press('tab'); pyautogui.press('tab')
                pyautogui.write('Item', interval=0.1); pyautogui.press('tab')
                pyautogui.write(cost, interval=0.1)
                pyautogui.press('tab'); pyautogui.press('tab')
                pyautogui.write(cost_element2, interval=0.1); pyautogui.press('tab')
                pyautogui.write(sub_element2, interval=0.1); pyautogui.press('tab')
                pyautogui.press('tab'); pyautogui.write('Total Value', interval=0.1)
                pyautogui.press('tab'); pyautogui.write(overhead, interval=0.1); pyautogui.press('tab')
                pyautogui.hotkey('ctrl', 's'); pyautogui.hotkey('ctrl', 'f4'); pyautogui.hotkey('ctrl', 'f4')
                ws[f'E{row_ix}'] = 'Inventory Asset Completed'
                print(item, 'Pending 존재 | 기본값 조정', 'sec:', round(time.time() - t0, 2))

    # ===== 2) Frozen(또는 Annual) 업데이트 루프 =====
    print('Frozen cost update')
    std_center = locate(IMG('standard.png'), timeout=10)
    if std_center:
        pyautogui.doubleClick(std_center)
        time.sleep(1)

        for row_ix, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            item_code, cost_type = fmt(row[0]), fmt(row[1])
            if cost_type not in ('Pending', 'Annual'):
                continue

            t0 = time.time()
            cost_remark = f"KRK_{item_code}_{dt.datetime.now():%Y%m%d%H%M}"

            upd_center = locate(IMG('nov_cost_update.png'), confidence=0.8, timeout=8)
            if not upd_center:
                print('NOV Update Costs 메뉴를 못 찾음')
                pyautogui.screenshot(cost_remark + '.png')
                break

            pyautogui.doubleClick(upd_center)
            pyautogui.press('tab', presses=2, interval=0.3)
            pyautogui.write(cost_type, interval=0.2); pyautogui.press('tab')
            pyautogui.write(cost_remark, interval=0.1)
            pyautogui.press('tab', presses=2, interval=0.3)
            pyautogui.write(item_code, interval=0.1)
            pyautogui.press('tab'); pyautogui.press('enter')
            pyautogui.press('tab', presses=3, interval=0.3); pyautogui.press('enter')

            summit = locate(IMG('summit_btn.png'), timeout=6)  # 파일명이 submit이면 이미지도/코드도 맞춰주세요
            if summit:
                pyautogui.click(summit)
                pyautogui.press('tab'); pyautogui.press('enter')
                ws[f'F{row_ix}'] = 'Cost update run'
                print(item_code, 'Cost Update sec:', round(time.time() - t0, 2))
    else:
        print('standard.png 메뉴 이미지 없음 – Frozen 업데이트 생략')

    print('END 전체 실행 시간:', round(time.time() - start_all, 2), '초')
    wb.save(EXCEL_OUT)
    pyautogui.alert('작업을 종료하였습니다!!!!')  # type: ignore

except Exception as e:
    print(e)
    pyautogui.alert(text=str(e), button='OK')  # type: ignore
    time.sleep(10)
