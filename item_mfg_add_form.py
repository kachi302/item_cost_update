import os
import time
import threading
import queue
import traceback
import datetime as dt
import re
import sys 
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk

import pyautogui as pag
from openpyxl import load_workbook


# =========================
# 경로/이미지 설정
# =========================
    

IMG_DIR = r"C:\item_cost_update\images"
CATEGORY_IMG = os.path.join(IMG_DIR, "ITEM_COST_CATEGORY.png")  # 화면에서 찾을 이미지
ITEM_ERROR = os.path.join(IMG_DIR, "error_pop.png")
SCREENSHOT_DIR = IMG_DIR                                       # 에러 스샷 저장 폴더



DEFAULT_SHEET = "Sheet1"
MFG_GROUP = "Item Cost Category"

# pyautogui 기본 안전장치
pag.FAILSAFE = True   # 마우스를 좌측상단(0,0)으로 이동하면 강제 중지(예외 발생)
pag.PAUSE = 0.03


# =========================
# 유틸 함수
# =========================
def safe_tab():
    """ctrl/alt/shift 눌림 꼬임 방지 후 TAB 한번"""
    for k in ("ctrl", "alt", "shift"):
        pag.keyUp(k)
    pag.keyDown("tab")
    time.sleep(0.5)
    pag.keyUp("tab")


def normalize_excel_text(val, to_upper=True) -> str:
    """
    엑셀 값 정규화:
      - 앞/뒤 공백만 제거 (strip)  ✅ 내부 공백은 유지
      - 내부의 탭/개행은 pyautogui에 위험하므로 스페이스로 치환 (공백 ' ' 자체는 유지)
      - (옵션) 대문자 변환
    """
    if val is None:
        return ""

    # 123.0 같은 값이 들어오면 '123'으로 정리
    if isinstance(val, float) and val.is_integer():
        s = str(int(val))
    else:
        s = str(val)

    # ✅ 시작/끝 공백만 제거
    s = s.strip()

    # ✅ 내부 공백(스페이스)은 그대로 두되, 탭/개행은 스페이스로 바꿔서 오작동 방지
    s = s.replace("\r", " ").replace("\n", " ").replace("\t", " ")

    if to_upper:
        s = s.upper()

    return s


class StopRequested(Exception):
    """사용자 Stop 버튼으로 중지 요청 시 사용"""
    pass


# =========================
# GUI 앱
# =========================
class ItemMfgGui(tk.Tk):
    def __init__(self):
        super().__init__()

        self.title("Item MFG 자동화 (Excel 선택 / 실행 / 중지)")
        self.geometry("950x650")

        self.excel_path = tk.StringVar(value="")
        self.status_var = tk.StringVar(value="대기 중")
        self.progress_text = tk.StringVar(value="진행: 0/0")
        self.progress_val = tk.IntVar(value=0)

        self.msg_q = queue.Queue()
        self.worker = None
        self.running = False
        self.stop_event = threading.Event()

        self._build_ui()
        # ✅ 엑셀 경로가 변경될 때마다 실행 버튼 상태 업데이트
        self.excel_path.trace_add("write", lambda *_: self._update_run_button_state())
        self._update_run_button_state()
        
        self._poll_queue()

        

    def _update_run_button_state(self):
    # 실행 중이면 무조건 비활성
        if self.running:
            self.btn_run.config(state="disabled")
            return

        path = self.excel_path.get().strip()
        valid = (
            bool(path)
            and os.path.isfile(path)
            and path.lower().endswith((".xlsx", ".xlsm"))
        )
        self.btn_run.config(state="normal" if valid else "disabled")

    
    # ---------- UI ----------
    def _build_ui(self):
        frm_top = ttk.Frame(self, padding=10)
        frm_top.pack(fill="x")

        ttk.Label(frm_top, text="엑셀 파일:").grid(row=0, column=0, sticky="w")
        ttk.Entry(frm_top, textvariable=self.excel_path, width=85).grid(row=0, column=1, sticky="we", padx=(6, 6))

        ttk.Button(frm_top, text="엑셀 선택", command=self.on_select_excel).grid(row=0, column=2, padx=(0, 6))
        self.btn_run = ttk.Button(frm_top, text="실행", command=self.on_run,state="dosabled")
        self.btn_run.grid(row=0, column=3, padx=(0, 6))
        self.btn_stop = ttk.Button(frm_top, text="실행중지", command=self.on_stop, state="disabled")
        self.btn_stop.grid(row=0, column=4)

        frm_top.columnconfigure(1, weight=1)

        frm_mid = ttk.Frame(self, padding=(10, 0, 10, 10))
        frm_mid.pack(fill="x")

        ttk.Label(frm_mid, textvariable=self.status_var).pack(anchor="w")
        self.pbar = ttk.Progressbar(frm_mid, orient="horizontal", mode="determinate",
                                    variable=self.progress_val, maximum=100)
        self.pbar.pack(fill="x", pady=(6, 2))
        ttk.Label(frm_mid, textvariable=self.progress_text).pack(anchor="w")

        frm_log = ttk.Frame(self, padding=10)
        frm_log.pack(fill="both", expand=True)

        # ✅ 로그 상단 바(라벨 + Clear 버튼)
        log_bar = ttk.Frame(frm_log)
        log_bar.pack(fill="x")

        ttk.Label(log_bar, text="실행 로그 (Item 단계별 진행):").pack(side="left")
        ttk.Button(log_bar, text="로그 Clear", command=self.on_clear_log).pack(side="right")

        self.txt = scrolledtext.ScrolledText(frm_log, wrap="word")
        self.txt.pack(fill="both", expand=True, pady=(6, 0))

        self._log("※ 사용 방법")
        self._log("  1) 엑셀 선택")
        self._log("  2) 대상 프로그램(ERP) 화면을 미리 활성화(포커스)")
        self._log("  3) 실행 클릭 → 3초 후 시작")
        self._log("  4) 중지하려면 '실행중지' 클릭 또는 마우스를 좌측상단(0,0)으로 이동(FAILSAFE)")
        self._log("")
        self._log(f"- 시트: {DEFAULT_SHEET}")
        self._log("- Excel 컬럼 사용: A=item, B=mfg_value")
        self._log("- 정규화: 앞/뒤 공백 제거(strip) + 대문자 변환 (내부 스페이스 유지)")
        self._log(f"- 그룹명 입력값: {MFG_GROUP}")
        self._log("- 이미지가 있으면(발견되면) → 해당 record는 SKIP 후 다음 record로 진행")
        self._log(f"- 화면 체크 이미지: {CATEGORY_IMG}")
        self._log("")

    def _log(self, msg: str):
        ts = dt.datetime.now().strftime("%H:%M:%S")
        self.txt.insert("end", f"[{ts}] {msg}\n")
        self.txt.see("end")

    def on_clear_log(self):
        """✅ 로그창 클리어 버튼"""
        self.txt.delete("1.0", "end")

    # ---------- 버튼 이벤트 ----------
    def on_select_excel(self):
        path = filedialog.askopenfilename(
            title="엑셀 파일 선택",
            filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")]
        )
        if path:
            self.excel_path.set(path)
            self._log(f"엑셀 선택: {path}")
            self._update_run_button_state()

    def on_run(self):
        if self.running:
            messagebox.showinfo("안내", "이미 실행 중입니다.")
            return

        xlsx = self.excel_path.get().strip()
        if not xlsx or not os.path.exists(xlsx):
            messagebox.showwarning("경고", "유효한 엑셀 파일을 선택하세요.")
            return

        ok = messagebox.askokcancel("확인", "Item MFG 추가 작업을 진행하시겠습니까?")
        if not ok:
            self._log("사용자가 실행을 취소했습니다.")
            return

        self.running = True
        self.stop_event.clear()
        self.btn_run.config(state="disabled")
        self.btn_stop.config(state="normal")
        self.status_var.set("실행 준비 중...")

        os.makedirs(SCREENSHOT_DIR, exist_ok=True)

        self.worker = threading.Thread(target=self._worker_run, args=(xlsx,), daemon=True)
        self.worker.start()

    def on_stop(self):
        if not self.running:
            return
        self.stop_event.set()
        self._log(">>> [STOP 요청] 실행중지 버튼이 눌렸습니다. 가능한 빠르게 안전 중지합니다...")

    # ---------- worker 유틸 ----------
    def _check_stop(self):
        if self.stop_event.is_set():
            raise StopRequested()

    def _sleep(self, seconds: float, step: float = 0.1):
        end = time.time() + seconds
        while time.time() < end:
            self._check_stop()
            time.sleep(min(step, end - time.time()))

    def _release_keys(self):
        for k in ("ctrl", "alt", "shift"):
            try:
                pag.keyUp(k)
            except Exception:
                pass

    # ---------- worker 본체 ----------
    def _worker_run(self, excel_path: str):
        try:
            self.msg_q.put(("STATUS", "엑셀 로딩 중..."))
            wb = load_workbook(excel_path, data_only=True)

            if DEFAULT_SHEET not in wb.sheetnames:
                raise ValueError(f"시트 '{DEFAULT_SHEET}'를 찾을 수 없습니다. 현재 시트: {wb.sheetnames}")

            ws = wb[DEFAULT_SHEET]

            max_row = ws.max_row or 0
            if max_row <= 0:
                raise ValueError("엑셀에 데이터가 없습니다.")

            total = max_row
            self.msg_q.put(("PROG_INIT", total))
            self.msg_q.put(("LOG", f"Start: {dt.datetime.now().strftime('%Y.%m%d - %H:%M:%S')}"))
            self.msg_q.put(("LOG", "3초 후 시작합니다. 대상 프로그램 화면을 활성화(포커스) 하세요."))
            self._sleep(3)

            rec = 0
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                self._check_stop()
                rec += 1
                self.msg_q.put(("PROG", rec, total))

                item_raw = row[0] if len(row) > 0 else None
                mfg_value_raw = row[1] if len(row) > 1 else None  # B열

                # ✅ 앞/뒤만 공백 제거 + 대문자(내부 스페이스 유지)
                item = normalize_excel_text(item_raw, to_upper=True)
                mfg_value = normalize_excel_text(mfg_value_raw, to_upper=True)

                if not item:
                    self.msg_q.put(("LOG", f"[SKIP] row={row_idx}: item 값이 비어있음"))
                    continue

                # ✅ 헤더 스킵은 공백 유무 상관 없이 판별되도록 별도 키 사용(내부 공백 제거는 '판별'에만 사용)
                header_key = re.sub(r"\s+", "", item)
                if header_key in ("ITEM", "ITEM_CODE", "ITEMCODE"):
                    self.msg_q.put(("LOG", f"[SKIP] row={row_idx}: 헤더로 판단되어 스킵 (item={item})"))
                    continue

                self.msg_q.put(("LOG", "----------------------------------------"))
                self.msg_q.put(("LOG", f"Rec: {rec}/{total} | row={row_idx} | item={item} | value(B)={mfg_value}"))
                self.msg_q.put(("STATUS", f"처리 중: {item}"))

                try:
                    self._sleep(1)

                    # 1) item 찾기
                    self.msg_q.put(("LOG", "  - F11"))
                    pag.hotkey("f11")
                    self._sleep(2)

                    self.msg_q.put(("LOG", "  - item 입력"))
                    pag.write(item, interval=0.08)
                    self._sleep(1)

                    self.msg_q.put(("LOG", "  - Ctrl+F11"))
                    pag.hotkey("ctrl", "f11")
                    self._sleep(2)
                    # item code 없는 경우
                    self.msg_q.put(("LOG", " - ITEM CODE 검색"))
                    # try:
                    #     err_png = pag.locateOnScreen(ITEM_ERROR, confidence=0.8)    
                    # except Exception as e:
                    #     self.msg_q.put("LOG", f"[SKIP] ITEM 존재 함 ")                                           
                        
                    # if err_png is not None:
                    #     self.msg_q.put(("LOG", f"  [WARN] ITEM 탐색 실패"))
                    #     err_png = None
                    #     pag.press("enter")
                    #     self._sleep(3)
                    #     pag.press("f4")
                    #     self._sleep(3)
                    #     self.msg_q.put(("LOG", f"[SKIP] item={item} (ITEM 없음, 다음 단계로 )"))
                    #     continue       
                    # 2) tools 진입
                    # ITEM_ERROR: error_pop.png 같은 에러 팝업 이미지 경로라고 가정

                    # 0) 파일 존재 여부부터 먼저 확인 (없으면 경고만 찍고 계속 진행)
                    if not os.path.isfile(ITEM_ERROR):
                        self.msg_q.put(("LOG", f"  [WARN] ITEM_ERROR 이미지 파일 없음 → 체크 없이 계속 진행: {ITEM_ERROR}"))
                        error_png = None
                    else:
                        # 1) 화면에서 에러 팝업 이미지 탐색
                        try:
                            error_png = pag.locateOnScreen(ITEM_ERROR, confidence=0.8)
                        except Exception as e:
                            self.msg_q.put(("LOG", f"  [WARN] ITEM_ERROR 이미지 탐색 중 예외 → 계속 진행: {e}"))
                            error_png = None

                    # 2) 이미지가 발견되면: Enter + 닫기 + SKIP
                    if error_png is not None:
                        self.msg_q.put(("LOG", "  - ITEM_ERROR 이미지 발견 → Enter 처리 후 화면 닫고 다음 record"))
                        pag.press("enter")
                        self._sleep(3)

                        # 닫기 키는 실제 시스템에 맞게 선택하세요(원 코드 유지: F4)
                        pag.press("f4")          # 원하신대로 F4
                        # pag.hotkey("ctrl","f4") # (대안) 창 닫기면 ctrl+f4
                        # pag.hotkey("alt","f4")  # (대안) alt+f4

                        self._sleep(3)
                        self.msg_q.put(("LOG", f"[SKIP] item={item} (ITEM 없음/에러 팝업)"))
                        continue

                    # 3) 여기로 오면 error_png == None → 계속 진행(정상 흐름)
                    self.msg_q.put(("LOG", "  - ITEM_ERROR 이미지 미발견 → 계속 진행"))

                    self.msg_q.put(("LOG", "  - Alt+T"))
                    pag.hotkey("alt", "t")
                    self._sleep(1)

                    self.msg_q.put(("LOG", "  - Enter (Tools 진입)"))
                    pag.press("enter")
                    self._sleep(3)

                    # 3) 이미지 체크: 이미지가 있으면 다음 record로 넘어가야 함
                    self.msg_q.put(("LOG", "  - 화면에서 ITEM_COST_CATEGORY 이미지 탐색"))
                    try:
                        category_png = pag.locateOnScreen(CATEGORY_IMG, confidence=0.8)
                    except Exception as img_e:
                        self.msg_q.put(("LOG", f"  [WARN] 이미지 탐색 실패: {img_e}"))
                        
                        category_png = None

                    if category_png is not None:
                        self.msg_q.put(("LOG", "  - 이미지 발견 → 해당 record SKIP, 닫고 다음 record로 이동"))
                        pag.hotkey("ctrl", "f4")
                        self._sleep(1.5)
                        self.msg_q.put(("LOG", f"[SKIP] item={item} (이미 존재로 판단)"))
                        continue

                    # 이미지가 없으면 입력/저장 진행
                    self.msg_q.put(("LOG", "  - 이미지 미발견 → 입력/저장 진행"))

                    self.msg_q.put(("LOG", "  - Ctrl+Down (마지막 라인 이동)"))
                    pag.hotkey("ctrl", "down")
                    self._sleep(2)

                    self.msg_q.put(("LOG", f"  - Group 입력: {MFG_GROUP}"))
                    pag.write(MFG_GROUP, interval=0.08)
                    self._sleep(2)

                    self.msg_q.put(("LOG", "  - TAB 이동"))
                    safe_tab()
                    self._sleep(2)

                    self.msg_q.put(("LOG", f"  - Value 입력: {mfg_value}"))
                    pag.write(mfg_value, interval=0.08)
                    self._sleep(2)

                    self.msg_q.put(("LOG", "  - TAB 이동"))
                    safe_tab()
                    self._sleep(1)

                    self.msg_q.put(("LOG", "  - Ctrl+S (저장)"))
                    pag.hotkey("ctrl", "s")
                    self._sleep(2)

                    self.msg_q.put(("LOG", "  - Ctrl+F4 (닫기)"))
                    pag.hotkey("ctrl", "f4")
                    self._sleep(1.5)

                    self.msg_q.put(("LOG", f"[OK] item={item} 완료"))

                except StopRequested:
                    raise
                except pag.FailSafeException:
                    self.msg_q.put(("LOG", "[FAILSAFE] 좌측상단 이동 감지로 중단합니다."))
                    raise StopRequested()
                except Exception as row_e:
                    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
                    safe_item = "".join(ch for ch in item if ch.isalnum() or ch in ("-", "_"))[:30]
                    shot = os.path.join(SCREENSHOT_DIR, f"error_row{row_idx}_{safe_item}_{ts}.png")
                    try:
                        pag.screenshot(shot)
                        self.msg_q.put(("LOG", f"[ERROR] row={row_idx}, item={item}"))
                        self.msg_q.put(("LOG", f"        스크린샷 저장: {shot}"))
                    except Exception:
                        self.msg_q.put(("LOG", f"[ERROR] row={row_idx}, item={item} (스크린샷 저장 실패 가능)"))

                    self.msg_q.put(("LOG", f"        예외: {row_e}"))
                    self.msg_q.put(("LOG", "        다음 row로 계속 진행합니다."))

            self.msg_q.put(("LOG", f"\nEnd: {dt.datetime.now().strftime('%Y.%m%d - %H:%M:%S')}"))
            self.msg_q.put(("DONE", "OK"))

        except StopRequested:
            self._release_keys()
            self.msg_q.put(("LOG", "\n[STOP] 사용자 요청으로 중지되었습니다."))
            self.msg_q.put(("DONE", "STOP"))

        except Exception as e:
            self._release_keys()
            self.msg_q.put(("LOG", "\n[치명적 오류 발생]"))
            self.msg_q.put(("LOG", str(e)))
            self.msg_q.put(("LOG", traceback.format_exc()))
            self.msg_q.put(("DONE", "ERROR"))

    # ---------- queue poll ----------
    def _poll_queue(self):
        try:
            while True:
                msg = self.msg_q.get_nowait()
                kind = msg[0]

                if kind == "LOG":
                    self._log(msg[1])

                elif kind == "STATUS":
                    self.status_var.set(msg[1])

                elif kind == "PROG_INIT":
                    total = max(int(msg[1]), 1)
                    self.pbar.config(maximum=total)
                    self.progress_val.set(0)
                    self.progress_text.set(f"진행: 0/{total}")

                elif kind == "PROG":
                    cur, total = int(msg[1]), int(msg[2])
                    self.progress_val.set(cur)
                    self.progress_text.set(f"진행: {cur}/{total}")

                elif kind == "DONE":
                    result = msg[1]
                    self.running = False
                    self.btn_run.config(state="normal")
                    self.btn_stop.config(state="disabled")
                    self._update_run_button_state()
                    self.status_var.set("대기 중")

                    if result == "OK":
                        messagebox.showinfo("완료", "작업을 종료합니다.")
                    elif result == "STOP":
                        messagebox.showinfo("중지", "사용자 요청으로 중지되었습니다.")
                    else:
                        messagebox.showerror("오류", "오류로 중단되었습니다. 로그를 확인하세요.")

        except queue.Empty:
            pass

        self.after(100, self._poll_queue)


if __name__ == "__main__":
    app = ItemMfgGui()
    app.mainloop()
