import os
import time
import datetime as dt
import pyautogui
from openpyxl import load_workbook

# 전역 설정
pyautogui.FAILSAFE = True         # 마우스 왼쪽 위로 이동 시 긴급 중지
pyautogui.PAUSE = 0.2             # 각 동작 사이 기본 대기

BASE = os.path.dirname(os.path.abspath(__file__))
EXCEL_IN  = os.path.join(BASE, 'data', 'pending.xlsx')
EXCEL_OUT = os.path.join(BASE, 'data', 'Frozen_JOB.xlsx')
IMG = lambda name: os.path.join(BASE, 'images', name)

def wait_and_center(image_path: str, timeout: float = 12):
    """이미지가 화면에 나타날 때까지 기다렸다가 중앙 좌표를 반환."""
    end = time.time() + timeout
    while time.time() < end:
        try:
            # opencv가 있으면 confidence 사용
            box = pyautogui.locateOnScreen(image_path, confidence=0.9)
        except TypeError:
            # opencv가 없으면 기본 방식
            box = pyautogui.locateOnScreen(image_path)
        if box:
            return pyautogui.center(box)
        time.sleep(0.3)
    return None

def main():
    print('Frozen Cost Update Process, 5초 후 시작됩니다. 해당 화면을 다른 모니터로 이동하세요.')
    time.sleep(5)

    print('Excel File 읽기')
    wb = load_workbook(EXCEL_IN, data_only=True)
    if 'Sheet1' not in wb.sheetnames:
        raise RuntimeError("엑셀에 'Sheet1' 시트가 없습니다.")
    ws = wb['Sheet1']

    start_all = time.time()
    job = pyautogui.confirm(text='Frozen Cost Update 작업을 진행하시겠습니까?', buttons=['OK', 'Cancel'])
    if job != 'OK':
        print('사용자가 취소했습니다.')
        return

    print('Frozen cost update 시작')

    std_center = wait_and_center(IMG('standard.png'), timeout=15)
    if not std_center:
        raise RuntimeError('standard.png 이미지를 화면에서 찾지 못했습니다.')
    pyautogui.doubleClick(std_center)
    time.sleep(0.5)

    rec = 0
    # 2행부터, 값만
    for row in ws.iter_rows(min_row=2, values_only=True):
        item_code = (str(row[0]).strip() if row and row[0] else '')
        if not item_code:
            continue

        rec += 1
        t0 = time.time()

        cost_type = 'Pending'
        cost_remark = f"KRK_{item_code}_{dt.datetime.now():%Y%m%d%H%M}"

        upd_center = wait_and_center(IMG('nov_cost_update.png'), timeout=10)
        if not upd_center:
            print('NOV Update Costs 메뉴를 찾지 못했습니다. 중단합니다.')
            break

        pyautogui.doubleClick(upd_center)
        time.sleep(0.5)

        pyautogui.press('tab', presses=2, interval=0.3)
        pyautogui.write(cost_type, interval=0.1)
        pyautogui.press('tab')
        pyautogui.write(cost_remark, interval=0.1)
        pyautogui.press('tab', presses=2, interval=0.3)
        pyautogui.write(item_code, interval=0.1)
        pyautogui.press('tab')
        pyautogui.press('enter')
        pyautogui.press('tab', presses=3, interval=0.25)
        pyautogui.press('enter')
        time.sleep(2)

        # 파일명이 summit_btn.png 인 듯하지만 submit 의미면 이미지 파일명도 맞춰주세요
        submit_center = wait_and_center(IMG('summit_btn.png'), timeout=10)
        if submit_center:
            pyautogui.click(submit_center)
            time.sleep(0.7)
            pyautogui.press('tab')
            pyautogui.press('enter')
            ws[f'F{rec}'] = 'Cost update run'
            print(item_code, '실행 시간:', round(time.time() - t0, 2), 's')
        else:
            print('Submit 버튼 이미지를 찾지 못했습니다. 다음 항목으로 넘어갑니다.')

    print('END 전체 실행 시간:', round(time.time() - start_all, 2), 's / 처리 건수:', rec)
    wb.save(EXCEL_OUT)
    pyautogui.alert('작업을 종료하였습니다!')

if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        pyautogui.alert(text=str(e), button='OK')
        # 필요 시 로그 확인을 위해 예외 재발생
        raise
